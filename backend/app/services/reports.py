"""Izvestaji o radu sluzbe, AI metrike i trosak, izvoz u Excel, Word i PDF."""
import io
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import AICall, AISuggestion, AuditLog, Case, Email, OrgUnit, PhoneContact
from . import catalog as cat

EDIT_LABELS = {"bez_izmena": "Bez izmena", "manje_izmene": "Manje izmene", "vece_izmene": "Veće izmene", "napisan_iznova": "Napisan iznova"}
FEATURE_LABELS = {"analiza": "Analiza obraćanja", "odgovor": "Predlog odgovora", "filter": "Filter pošte",
                  "mreze_tekst": "Tekst za mreže", "mreze_skracivanje": "Skraćivanje teksta", "pozadina_slika": "AI pozadina (slika)"}


def _now():
    return datetime.now(timezone.utc)


def period(days: int) -> tuple[datetime, datetime]:
    end = _now()
    return end - timedelta(days=days), end


def summary(db: Session, days: int = 30) -> dict:
    start, end = period(days)
    labels = {c["code"]: c["label"] for c in cat.categories(db)}
    units = {u.id: u for u in db.scalars(select(OrgUnit)).all()}
    cases = db.scalars(select(Case).where(Case.created_at >= start)).all()
    open_cases = db.scalars(select(Case).where(Case.status.in_(["novo", "u_obradi", "ceka_proveru"]))).all()
    closed = [c for c in cases if c.status == "zatvoreno"]
    responded = [c for c in cases if c.first_response_at]
    resp_hours = [(c.first_response_at - c.created_at).total_seconds() / 3600 for c in responded]
    on_time = [c for c in responded if c.due_at and c.first_response_at <= c.due_at]
    filtered = db.scalar(select(func.count()).select_from(Email).where(Email.received_at >= start, Email.folder.in_(["spam", "tehnicka"])))
    calls = db.scalar(select(func.count()).select_from(PhoneContact).where(PhoneContact.started_at >= start))
    overdue = [c for c in open_cases if c.due_at and c.due_at < _now()]

    by_day = defaultdict(lambda: {"primljeno": 0, "zatvoreno": 0})
    for i in range(days + 1):
        d = (start + timedelta(days=i)).date().isoformat()
        by_day[d]
    for c in cases:
        by_day[c.created_at.date().isoformat()]["primljeno"] += 1
    for c in db.scalars(select(Case).where(Case.closed_at >= start, Case.closed_at <= end)).all():
        by_day[c.closed_at.date().isoformat()]["zatvoreno"] += 1

    by_cat = Counter(c.category for c in cases)
    unit_rows = []
    for uid, u in units.items():
        uc = [c for c in cases if c.org_unit_id == uid]
        ur = [(c.first_response_at - c.created_at).total_seconds() / 3600 for c in uc if c.first_response_at]
        unit_rows.append({"jedinica": u.code, "naziv": u.name, "primljeno": len(uc),
                          "zatvoreno": len([c for c in uc if c.status == "zatvoreno"]),
                          "otvoreno_sada": len([c for c in open_cases if c.org_unit_id == uid]),
                          "prekoraceno_sada": len([c for c in overdue if c.org_unit_id == uid]),
                          "prosek_odgovora_h": round(sum(ur) / len(ur), 1) if ur else None})
    return {
        "period": {"od": start, "do": end, "dana": days},
        "kpi": {
            "primljeno": len(cases), "zatvoreno": len(closed), "otvoreno_sada": len(open_cases), "prekoraceno_sada": len(overdue),
            "prosek_prvog_odgovora_h": round(sum(resp_hours) / len(resp_hours), 1) if resp_hours else None,
            "u_roku_procenat": round(100 * len(on_time) / len(responded), 1) if responded else None,
            "odvojena_posta": filtered, "telefonski_kontakti": calls,
        },
        "po_danima": [{"datum": d, **v} for d, v in sorted(by_day.items())],
        "po_kategorijama": [{"kategorija": labels.get(k, k), "broj": n} for k, n in by_cat.most_common()],
        "po_kanalima": [{"kanal": k, "broj": n} for k, n in Counter(c.channel for c in cases).most_common()],
        "po_jedinicama": sorted(unit_rows, key=lambda r: r["primljeno"], reverse=True),
        "otvoreni_po_statusu": [{"status": k, "broj": n} for k, n in Counter(c.status for c in open_cases).most_common()],
    }


def ai_metrics(db: Session, days: int = 30) -> dict:
    start, _ = period(days)
    units = {u.id: u.code for u in db.scalars(select(OrgUnit)).all()}
    labels = {c["code"]: c["label"] for c in cat.categories(db)}
    sugs = db.scalars(select(AISuggestion).where(AISuggestion.created_at >= start)).all()
    fin_unit = [s for s in sugs if s.accepted_unit is not None]
    fin_text = [s for s in sugs if s.edit_class]
    edit = Counter(s.edit_class for s in fin_text)
    confusion = Counter((units.get(s.suggested_unit_id), units.get(s.final_unit_id)) for s in fin_unit)
    per_cat = defaultdict(lambda: {"ukupno": 0, "jedinica_ok": 0, "bez_vecih_izmena": 0, "sa_tekstom": 0})
    for s in fin_unit:
        r = per_cat[s.category]
        r["ukupno"] += 1
        r["jedinica_ok"] += bool(s.accepted_unit)
    for s in fin_text:
        r = per_cat[s.category]
        r["sa_tekstom"] += 1
        r["bez_vecih_izmena"] += s.edit_class in ("bez_izmena", "manje_izmene")
    weekly = defaultdict(lambda: {"ukupno": 0, "korisceno": 0})
    for s in fin_text:
        wk = (s.finalized_at or s.created_at).strftime("%G-N%V")
        weekly[wk]["ukupno"] += 1
        weekly[wk]["korisceno"] += s.edit_class in ("bez_izmena", "manje_izmene")
    tpl_use = Counter(src["code"] for s in sugs for src in (s.sources or []) if src.get("koriscen"))
    return {
        "predloga": len(sugs), "potvrdjeno_rasporeda": len(fin_unit), "zavrseno_odgovora": len(fin_text),
        "tacnost_jedinice": round(100 * sum(bool(s.accepted_unit) for s in fin_unit) / len(fin_unit), 1) if fin_unit else None,
        "tacnost_prioriteta": round(100 * sum(bool(s.accepted_priority) for s in fin_unit) / len(fin_unit), 1) if fin_unit else None,
        "izmene": [{"klasa": k, "naziv": EDIT_LABELS[k], "broj": edit.get(k, 0),
                    "procenat": round(100 * edit.get(k, 0) / len(fin_text), 1) if fin_text else 0} for k in EDIT_LABELS],
        "prosecna_izmena": round(sum(s.edit_ratio or 0 for s in fin_text) / len(fin_text), 3) if fin_text else None,
        "matrica": [{"predlog": a, "konacno": b, "broj": n} for (a, b), n in sorted(confusion.items(), key=lambda x: -x[1])],
        "po_kategorijama": sorted([{"kategorija": labels.get(k, k), **v,
                                    "jedinica_procenat": round(100 * v["jedinica_ok"] / v["ukupno"], 1) if v["ukupno"] else None,
                                    "tekst_procenat": round(100 * v["bez_vecih_izmena"] / v["sa_tekstom"], 1) if v["sa_tekstom"] else None}
                                   for k, v in per_cat.items()], key=lambda r: -r["ukupno"]),
        "nedeljno": [{"nedelja": k, **v, "procenat": round(100 * v["korisceno"] / v["ukupno"], 1) if v["ukupno"] else 0} for k, v in sorted(weekly.items())],
        "obrasci": [{"sifra": k, "broj": n} for k, n in tpl_use.most_common(10)],
    }


def ai_cost(db: Session, days: int = 30) -> dict:
    start, _ = period(days)
    price_in = (cat.setting(db, "ai_price_input_per_1m", {}) or {}).get("value", 0)
    price_out = (cat.setting(db, "ai_price_output_per_1m", {}) or {}).get("value", 0)
    price_img = (cat.setting(db, "ai_image_price", {}) or {}).get("value", 0)
    calls = db.scalars(select(AICall).where(AICall.at >= start)).all()

    def cost(c: AICall) -> float:
        if c.feature == "pozadina_slika":
            return price_img if c.ok else 0
        return (c.prompt_tokens * price_in + c.completion_tokens * price_out) / 1_000_000

    feat = defaultdict(lambda: {"poziva": 0, "ulaz": 0, "izlaz": 0, "trosak": 0.0, "lat": 0, "greske": 0})
    day = defaultdict(lambda: {"ulaz": 0, "izlaz": 0, "trosak": 0.0})
    modes = Counter()
    for c in calls:
        f = feat[c.feature]
        f["poziva"] += 1
        f["ulaz"] += c.prompt_tokens
        f["izlaz"] += c.completion_tokens
        f["trosak"] += cost(c)
        f["lat"] += c.latency_ms
        f["greske"] += not c.ok
        d = day[c.at.date().isoformat()]
        d["ulaz"] += c.prompt_tokens
        d["izlaz"] += c.completion_tokens
        d["trosak"] += cost(c)
        modes[c.mode] += 1
    total = sum(cost(c) for c in calls)
    analyzed = len({c.case_id for c in calls if c.case_id})
    return {
        "cene": {"ulaz_1m": price_in, "izlaz_1m": price_out, "slika": price_img, "valuta": "RSD",
                 "napomena": "Interna cena: amortizacija servera, struja i održavanje podeljeni sa očekivanim brojem tokena. Podešava se u administraciji."},
        "ukupno": {"poziva": len(calls), "ulaz": sum(c.prompt_tokens for c in calls), "izlaz": sum(c.completion_tokens for c in calls),
                   "trosak": round(total, 2), "po_predmetu": round(total / analyzed, 3) if analyzed else None, "predmeta": analyzed,
                   "rezervni_rezim": modes.get("mock", 0)},
        "po_funkciji": [{"funkcija": k, "naziv": FEATURE_LABELS.get(k, k), "poziva": v["poziva"], "ulaz": v["ulaz"], "izlaz": v["izlaz"],
                         "trosak": round(v["trosak"], 2), "prosecno_ms": int(v["lat"] / v["poziva"]) if v["poziva"] else 0, "greske": v["greske"]}
                        for k, v in sorted(feat.items(), key=lambda x: -x[1]["trosak"])],
        "po_danima": [{"datum": k, **{kk: (round(vv, 2) if isinstance(vv, float) else vv) for kk, vv in v.items()}} for k, v in sorted(day.items())],
    }


def audit(db: Session, limit: int = 300) -> list[dict]:
    rows = db.scalars(select(AuditLog).order_by(AuditLog.at.desc()).limit(limit)).all()
    return [{"id": r.id, "at": r.at, "actor": r.actor, "entity": r.entity, "entity_id": r.entity_id, "action": r.action,
             "summary": r.summary, "diff": r.diff} for r in rows]


# ------------------------------------------------------------------ izvoz

def _fmt(v):
    if v is None:
        return "-"
    if isinstance(v, float):
        return f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


def export_xlsx(db: Session, days: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    s, a, c = summary(db, days), ai_metrics(db, days), ai_cost(db, days)
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0B1F3A")

    def sheet(title, headers, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head, fill
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(48, max(len(str(x.value or "")) for x in col) + 2))
        ws.freeze_panes = "A2"
        return ws

    k = s["kpi"]
    sheet("Pregled", ["Pokazatelj", "Vrednost"], [
        ["Period (dana)", days], ["Primljeno predmeta", k["primljeno"]], ["Zatvoreno", k["zatvoreno"]],
        ["Otvoreno sada", k["otvoreno_sada"]], ["Prekoračen rok (sada)", k["prekoraceno_sada"]],
        ["Prosečno vreme do prvog odgovora (h)", k["prosek_prvog_odgovora_h"]], ["Odgovoreno u roku (%)", k["u_roku_procenat"]],
        ["Odvojena pošta (spam i tehnička)", k["odvojena_posta"]], ["Telefonski kontakti", k["telefonski_kontakti"]],
        ["Tačnost AI predloga jedinice (%)", a["tacnost_jedinice"]], ["Tačnost AI predloga prioriteta (%)", a["tacnost_prioriteta"]],
        ["AI trošak (RSD, interna cena)", c["ukupno"]["trosak"]],
    ], first=True)
    sheet("Po danima", ["Datum", "Primljeno", "Zatvoreno"], [[r["datum"], r["primljeno"], r["zatvoreno"]] for r in s["po_danima"]])
    sheet("Po kategorijama", ["Kategorija", "Broj"], [[r["kategorija"], r["broj"]] for r in s["po_kategorijama"]])
    sheet("Po jedinicama", ["Jedinica", "Naziv", "Primljeno", "Zatvoreno", "Otvoreno sada", "Prekoračeno sada", "Prosek odgovora (h)"],
          [[r["jedinica"], r["naziv"], r["primljeno"], r["zatvoreno"], r["otvoreno_sada"], r["prekoraceno_sada"], r["prosek_odgovora_h"]] for r in s["po_jedinicama"]])
    sheet("AI predlozi", ["Klasa izmene", "Broj", "Procenat"], [[r["naziv"], r["broj"], r["procenat"]] for r in a["izmene"]])
    sheet("AI po kategorijama", ["Kategorija", "Predloga", "Tačna jedinica (%)", "Bez većih izmena (%)"],
          [[r["kategorija"], r["ukupno"], r["jedinica_procenat"], r["tekst_procenat"]] for r in a["po_kategorijama"]])
    sheet("AI trošak", ["Funkcija", "Poziva", "Ulazni tokeni", "Izlazni tokeni", "Trošak (RSD)", "Prosečno trajanje (ms)"],
          [[r["naziv"], r["poziva"], r["ulaz"], r["izlaz"], r["trosak"], r["prosecno_ms"]] for r in c["po_funkciji"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_docx(db: Session, days: int) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    s, a, c = summary(db, days), ai_metrics(db, days), ai_cost(db, days)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    doc.add_heading("Izveštaj o radu Korisničkog centra", 0)
    p = doc.add_paragraph(f"Period: poslednjih {days} dana (do {_now():%d.%m.%Y.}). Demo podaci, izmišljeni.")
    p.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = _fmt(v)
        doc.add_paragraph()

    k = s["kpi"]
    doc.add_heading("Ključni pokazatelji", 1)
    table(["Pokazatelj", "Vrednost"], [["Primljeno predmeta", k["primljeno"]], ["Zatvoreno", k["zatvoreno"]],
                                      ["Otvoreno sada", k["otvoreno_sada"]], ["Prekoračen rok (sada)", k["prekoraceno_sada"]],
                                      ["Prosečno vreme do prvog odgovora (h)", k["prosek_prvog_odgovora_h"]],
                                      ["Odgovoreno u roku (%)", k["u_roku_procenat"]], ["Odvojena pošta", k["odvojena_posta"]],
                                      ["Telefonski kontakti", k["telefonski_kontakti"]]])
    doc.add_heading("Rad po jedinicama", 1)
    table(["Jedinica", "Primljeno", "Zatvoreno", "Otvoreno", "Prekoračeno", "Prosek odgovora (h)"],
          [[r["jedinica"], r["primljeno"], r["zatvoreno"], r["otvoreno_sada"], r["prekoraceno_sada"], r["prosek_odgovora_h"]] for r in s["po_jedinicama"]])
    doc.add_heading("Teme obraćanja", 1)
    table(["Kategorija", "Broj"], [[r["kategorija"], r["broj"]] for r in s["po_kategorijama"][:12]])
    doc.add_heading("Podrška veštačke inteligencije", 1)
    doc.add_paragraph(f"Tačnost predloga nadležne jedinice: {_fmt(a['tacnost_jedinice'])}%. Tačnost predloga prioriteta: {_fmt(a['tacnost_prioriteta'])}%.")
    table(["Korišćenje predloga odgovora", "Broj", "Procenat"], [[r["naziv"], r["broj"], r["procenat"]] for r in a["izmene"]])
    doc.add_heading("Trošak veštačke inteligencije", 1)
    doc.add_paragraph(f"Ukupno {c['ukupno']['poziva']} poziva, {c['ukupno']['ulaz']:,} ulaznih i {c['ukupno']['izlaz']:,} izlaznih tokena. "
                      f"Interni trošak: {_fmt(c['ukupno']['trosak'])} RSD ({_fmt(c['ukupno']['po_predmetu'])} RSD po predmetu). {c['cene']['napomena']}".replace(",", "."))
    table(["Funkcija", "Poziva", "Trošak (RSD)", "Prosečno (ms)"], [[r["naziv"], r["poziva"], r["trosak"], r["prosecno_ms"]] for r in c["po_funkciji"]])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(db: Session, days: int) -> bytes:
    from jinja2 import Environment, FileSystemLoader
    from playwright.sync_api import sync_playwright

    s, a, c = summary(db, days), ai_metrics(db, days), ai_cost(db, days)
    env = Environment(loader=FileSystemLoader(Path(__file__).parent.parent / "social" / "templates"))
    max_day = max([r["primljeno"] for r in s["po_danima"]] + [1])
    html = env.get_template("report.html").render(s=s, a=a, c=c, days=days, now=_now(), fmt=_fmt, max_day=max_day,
                                                  font_dir=(Path(__file__).parent.parent / "social" / "assets" / "fonts").resolve().as_uri())
    import tempfile

    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as tmp:
        tmp.write(html)
    with sync_playwright() as p:
        b = p.chromium.launch(args=["--no-sandbox", "--allow-file-access-from-files"])
        page = b.new_page()
        page.goto(Path(tmp.name).as_uri(), wait_until="networkidle")
        pdf = page.pdf(format="A4", print_background=True, margin={"top": "16mm", "bottom": "16mm", "left": "14mm", "right": "14mm"})
        b.close()
    Path(tmp.name).unlink(missing_ok=True)
    return pdf
